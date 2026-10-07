import pandas as pd
import docx
import re
# Actas_Mesa_Operativa.py
# Lectura y entendimiento del archivo Actas_Mesa_Operativa.docx

# ---- PASO 1: abrir el documento y sacar sus párrafos ----
# Document abre el Word; cada párrafo tiene su texto en .text
documento = docx.Document("Actas_Mesa_Operativa.docx")
parrafos = [p.text.strip() for p in documento.paragraphs]

print("PASO 1 - Lectura del documento")
print("Párrafos:", len(parrafos))
print("Tablas dentro del documento:", len(documento.tables))

# ---- PASO 2: revisar el comienzo para identificar el patrón ----
print()
print("PASO 2 - Primeros 18 párrafos")
for numero, texto in enumerate(parrafos[:18]):
    print("   ", numero, "->", texto)

# ---- PASO 3: recorrer el documento y extraer los compromisos ----
# Patrón de la línea de compromiso. Cada (.*?) captura un dato:
# número. tema. Responsable: nombre. Fecha objetivo: fecha. Estado reportado: estado.
patron = re.compile(
    r"^(\d+)\. (.*?)\. Responsable: (.*?)\. Fecha objetivo: (.*?)\. Estado reportado: (.*?)\.$"
)

compromisos = []
sesion = None
fecha_sesion = None
area = None

for texto in parrafos:
    # Encabezado: empieza una nueva acta
    if texto.startswith("Mesa Operativa"):
        sesion = int(texto.split("Sesión")[1])
    elif texto.startswith("Fecha:"):
        fecha_sesion = texto.split(":", 1)[1].strip()
    elif texto.startswith("Área convocante:"):
        area = texto.split(":", 1)[1].strip()
    else:
        # Si el párrafo cumple el patrón, es un compromiso
        encontrado = patron.match(texto)
        if encontrado:
            numero, tema, responsable, fecha_objetivo, estado = encontrado.groups()
            compromisos.append(
                {
                    "Sesion": sesion,
                    "FechaSesion": fecha_sesion,
                    "Area": area,
                    "Numero": int(numero),
                    "Compromiso": tema,
                    "Responsable": responsable,
                    "FechaObjetivo": fecha_objetivo,
                    "Estado": estado,
                }
            )

actas = pd.DataFrame(compromisos)

# Las fechas vienen como día/mes/año
actas["FechaSesion"] = pd.to_datetime(actas["FechaSesion"], format="%d/%m/%Y")
actas["FechaObjetivo"] = pd.to_datetime(actas["FechaObjetivo"], format="%d/%m/%Y")

print()
print("PASO 3 - Compromisos extraídos")
print("Actas (sesiones):", actas["Sesion"].nunique())
print("Compromisos:", len(actas))
print("Compromisos por acta:", actas.groupby("Sesion").size().unique())
print(
    "Periodo:",
    actas["FechaSesion"].min().date(),
    "a",
    actas["FechaSesion"].max().date(),
)
print()
print("Primer compromiso:")
for campo, valor in actas.iloc[0].items():
    print("   ", campo, "=", valor)

# ---- PASO 4: estados, temas, responsables y vencimientos ----
print()
print("PASO 4 - Análisis de los compromisos")

print()
print("Estado reportado:")
print(actas["Estado"].value_counts())

print()
print("Temas distintos:", actas["Compromiso"].nunique())
print(actas["Compromiso"].value_counts())

# Responsables contra el catálogo oficial
resp = pd.read_excel("Catalogos_Gobierno.xlsx", sheet_name="Responsables")
en_catalogo = actas["Responsable"].isin(resp["Responsable"])
print()
print("Responsables distintos en las actas:", actas["Responsable"].nunique())
print(
    "Compromisos cuyo responsable existe en el catálogo:",
    en_catalogo.sum(),
    "de",
    len(actas),
)

# Abiertos y vencidos, medidos contra la fecha de la última sesión
corte = actas["FechaSesion"].max()
actas["Abierto"] = actas["Estado"] != "Cerrado"
actas["Vencido"] = actas["Abierto"] & (actas["FechaObjetivo"] < corte)
print()
print("Compromisos abiertos:", actas["Abierto"].sum())
print("Abiertos y con la fecha objetivo vencida:", actas["Vencido"].sum())

# Reabierto: el mismo tema aparece abierto después de haberse reportado cerrado
actas = actas.sort_values(["Compromiso", "Sesion"])
estado_anterior = actas.groupby("Compromiso")["Estado"].shift()
actas["Reabierto"] = actas["Abierto"] & (estado_anterior == "Cerrado")
print("Reaparecen abiertos después de reportarse cerrados:", actas["Reabierto"].sum())

actas.sort_values(["Sesion", "Numero"]).to_csv(
    "Actas_limpia.csv", index=False, encoding="utf-8-sig"
)
print()
print("Archivo guardado: Actas_limpia.csv")